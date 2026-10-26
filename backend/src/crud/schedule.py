"""排课占档 CRUD。"""
from __future__ import annotations

from datetime import date, datetime, time, timedelta
from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import Session

from model.lab import Lab
from model.schedule import LabScheduleBlock


async def list_blocks(
    db: AsyncSession,
    *,
    lab_id: int | None,
    start: datetime,
    end: datetime,
) -> list[LabScheduleBlock]:
    stmt = select(LabScheduleBlock).where(
        LabScheduleBlock.start_time < end,
        LabScheduleBlock.end_time > start,
    )
    if lab_id is not None:
        stmt = stmt.where(LabScheduleBlock.lab_id == lab_id)
    stmt = stmt.order_by(LabScheduleBlock.start_time.asc())
    return list((await db.execute(stmt)).scalars().all())


async def get_block(db: AsyncSession, block_id: int) -> LabScheduleBlock | None:
    return await db.get(LabScheduleBlock, block_id)


async def create_block(
    db: AsyncSession,
    *,
    lab_id: int,
    title: str,
    start_time: datetime,
    end_time: datetime,
    source: str,
    created_by: int | None,
    note: str | None,
) -> LabScheduleBlock:
    row = LabScheduleBlock(
        lab_id=lab_id,
        title=title,
        start_time=start_time,
        end_time=end_time,
        source=source,
        created_by=created_by,
        note=note,
    )
    db.add(row)
    await db.flush()
    await db.commit()
    await db.refresh(row)
    return row


async def delete_block(db: AsyncSession, row: LabScheduleBlock) -> None:
    await db.delete(row)
    await db.commit()


async def bulk_create(
    db: AsyncSession,
    rows: list[dict],
) -> int:
    for r in rows:
        db.add(LabScheduleBlock(**r))
    await db.flush()
    await db.commit()
    return len(rows)


def build_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "周循环"
    ws.append(
        [
            "实验室名称",
            "课程名",
            "星期",
            "开始时刻",
            "结束时刻",
            "学期开始日",
            "学期结束日",
            "备注",
        ]
    )
    ws.append(
        [
            "网络实验室",
            "计网实验",
            2,
            "08:00",
            "09:40",
            "2026-09-01",
            "2026-12-28",
            "计科2301",
        ]
    )
    ws2 = wb.create_sheet("绝对时间")
    ws2.append(["实验室名称", "课程名", "开始", "结束", "备注"])
    ws2.append(
        [
            "嵌入式实验室 A",
            "期中考试",
            "2026-11-10 14:00",
            "2026-11-10 16:00",
            "临时占档",
        ]
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip().replace("/", "-")
    return datetime.strptime(s[:10], "%Y-%m-%d").date()



def _parse_time(v) -> time:
    if isinstance(v, datetime):
        return v.time().replace(second=0, microsecond=0)
    if hasattr(v, "hour"):
        return time(v.hour, v.minute, getattr(v, "second", 0) or 0)
    s = str(v).strip()
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(s, fmt).time()
        except ValueError:
            continue
    raise ValueError(f"时刻格式错误: {v}")


def _parse_dt(v) -> datetime:
    if isinstance(v, datetime):
        return v.replace(microsecond=0)
    s = str(v).strip().replace("T", " ").replace("/", "-")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    raise ValueError(f"日期时间格式错误: {v}")


def _lab_map(session: Session) -> dict[str, int]:
    rows = list(session.execute(select(Lab)).scalars().all())
    return {r.name.strip(): r.id for r in rows}


def expand_weekly(
    *,
    lab_id: int,
    title: str,
    weekday: int,
    start_t: time,
    end_t: time,
    term_start,
    term_end,
    note: str | None,
    source: str,
    created_by: int | None,
) -> list[dict]:
    """weekday: 1=周一 … 7=周日。"""
    if weekday < 1 or weekday > 7:
        raise ValueError("星期须为 1～7（周一=1）")
    if end_t <= start_t:
        raise ValueError("结束时刻须晚于开始时刻")
    if term_end < term_start:
        raise ValueError("学期结束日须不早于开始日")

    out: list[dict] = []
    d = term_start
    # Python: Monday=0 … Sunday=6；我们用 1=周一 → weekday-1
    target = weekday - 1
    while d <= term_end:
        if d.weekday() == target:
            st = datetime.combine(d, start_t)
            et = datetime.combine(d, end_t)
            out.append(
                {
                    "lab_id": lab_id,
                    "title": title,
                    "start_time": st,
                    "end_time": et,
                    "source": source,
                    "created_by": created_by,
                    "note": note,
                }
            )
        d += timedelta(days=1)
    return out


def parse_excel_preview(
    session: Session,
    content: bytes,
    *,
    created_by: int | None,
    source: str = "excel",
) -> dict:
    """解析模板，返回 ok_rows / errors / blocks（可直接入库的字段）。"""
    labs = _lab_map(session)
    wb = load_workbook(BytesIO(content), data_only=True)
    errors: list[dict] = []
    blocks: list[dict] = []
    ok_rows = 0

    if "周循环" in wb.sheetnames:
        ws = wb["周循环"]
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue
            try:
                lab_name = str(row[0]).strip()
                title = str(row[1]).strip()
                weekday = int(row[2])
                start_t = _parse_time(row[3])
                end_t = _parse_time(row[4])
                term_start = _parse_date(row[5])
                term_end = _parse_date(row[6])
                note = str(row[7]).strip() if len(row) > 7 and row[7] else None
                if lab_name not in labs:
                    raise ValueError(f"实验室不存在: {lab_name}")
                expanded = expand_weekly(
                    lab_id=labs[lab_name],
                    title=title,
                    weekday=weekday,
                    start_t=start_t,
                    end_t=end_t,
                    term_start=term_start,
                    term_end=term_end,
                    note=note,
                    source=source,
                    created_by=created_by,
                )
                if not expanded:
                    raise ValueError("学期内无匹配该星期的日期")
                blocks.extend(expanded)
                ok_rows += 1
            except Exception as e:
                errors.append({"sheet": "周循环", "row": i, "error": str(e)})

    if "绝对时间" in wb.sheetnames:
        ws = wb["绝对时间"]
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue
            try:
                lab_name = str(row[0]).strip()
                title = str(row[1]).strip()
                st = _parse_dt(row[2])
                et = _parse_dt(row[3])
                note = str(row[4]).strip() if len(row) > 4 and row[4] else None
                if lab_name not in labs:
                    raise ValueError(f"实验室不存在: {lab_name}")
                if et <= st:
                    raise ValueError("结束须晚于开始")
                blocks.append(
                    {
                        "lab_id": labs[lab_name],
                        "title": title,
                        "start_time": st,
                        "end_time": et,
                        "source": source,
                        "created_by": created_by,
                        "note": note,
                    }
                )
                ok_rows += 1
            except Exception as e:
                errors.append({"sheet": "绝对时间", "row": i, "error": str(e)})

    # 与库内已有占档、本批次内部重叠 → 跳过（不整批失败）
    accepted: list[dict] = []
    skipped: list[dict] = []
    if blocks:
        min_st = min(b["start_time"] for b in blocks)
        max_et = max(b["end_time"] for b in blocks)
        existing = list(
            session.execute(
                select(LabScheduleBlock).where(
                    LabScheduleBlock.start_time < max_et,
                    LabScheduleBlock.end_time > min_st,
                )
            ).scalars().all()
        )

        def _overlap(a0, a1, b0, b1) -> bool:
            return a0 < b1 and a1 > b0

        for b in blocks:
            hit = None
            for e in existing:
                if e.lab_id == b["lab_id"] and _overlap(
                    b["start_time"], b["end_time"], e.start_time, e.end_time
                ):
                    hit = f"与已有《{e.title}》{e.start_time}重叠"
                    break
            if not hit:
                for a in accepted:
                    if a["lab_id"] == b["lab_id"] and _overlap(
                        b["start_time"],
                        b["end_time"],
                        a["start_time"],
                        a["end_time"],
                    ):
                        hit = f"与本批《{a['title']}》{a['start_time']}重叠"
                        break
            if hit:
                skipped.append(
                    {
                        "lab_id": b["lab_id"],
                        "title": b["title"],
                        "start_time": str(b["start_time"]),
                        "reason": hit,
                    }
                )
            else:
                accepted.append(b)

    return {
        "ok_rows": ok_rows,
        "block_count": len(accepted),
        "skipped_count": len(skipped),
        "skipped": skipped[:30],
        "errors": errors,
        "blocks": accepted,
    }



def sync_list_blocks(
    session: Session,
    *,
    lab_name: str | None,
    start: datetime | None,
    end: datetime | None,
) -> list[dict]:
    stmt = select(LabScheduleBlock)
    if lab_name:
        labs = _lab_map(session)
        lid = labs.get(lab_name.strip())
        if lid is None:
            return []
        stmt = stmt.where(LabScheduleBlock.lab_id == lid)
    if start:
        stmt = stmt.where(LabScheduleBlock.end_time > start)
    if end:
        stmt = stmt.where(LabScheduleBlock.start_time < end)
    stmt = stmt.order_by(LabScheduleBlock.start_time.asc()).limit(100)
    rows = list(session.execute(stmt).scalars().all())
    lab_by_id = {r.id: r.name for r in session.execute(select(Lab)).scalars().all()}
    return [
        {
            "id": r.id,
            "lab_id": r.lab_id,
            "lab_name": lab_by_id.get(r.lab_id),
            "title": r.title,
            "start_time": str(r.start_time),
            "end_time": str(r.end_time),
            "note": r.note,
            "source": r.source,
        }
        for r in rows
    ]


def sync_add_block(
    session: Session,
    *,
    user_id: int,
    lab_name: str,
    title: str,
    start_time: datetime,
    end_time: datetime,
    note: str | None = None,
) -> dict:
    labs = _lab_map(session)
    lid = labs.get(lab_name.strip())
    if lid is None:
        return {"ok": False, "error": f"实验室不存在: {lab_name}"}
    if end_time <= start_time:
        return {"ok": False, "error": "结束时间必须晚于开始时间"}
    # 与已有占档简单重叠检查
    overlap = session.execute(
        select(LabScheduleBlock).where(
            LabScheduleBlock.lab_id == lid,
            LabScheduleBlock.start_time < end_time,
            LabScheduleBlock.end_time > start_time,
        )
    ).scalars().first()
    if overlap:
        return {
            "ok": False,
            "error": f"与已有排课《{overlap.title}》重叠",
        }
    row = LabScheduleBlock(
        lab_id=lid,
        title=title.strip(),
        start_time=start_time,
        end_time=end_time,
        source="agent",
        created_by=user_id,
        note=note,
    )
    session.add(row)
    session.flush()
    session.refresh(row)
    return {"ok": True, "id": row.id, "title": row.title}
